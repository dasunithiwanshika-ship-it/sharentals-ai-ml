"""Excel workbook generator.

Generates the multi-sheet output workbook with:
  Sheet 1: Paste_eToll_Data  — definitive Master sheet (exact template match)
  Sheet 2: New Tolls          — only new chargeable records this run
  Sheet 3: Customer Summary   — grouped totals for 365 invoicing reference
  Sheet 4: Duplicates         — LINKT records skipped as duplicates
  Sheet 5: Unmatched          — tolls requiring manual driver review
  Sheet 6: Already Paid       — tolls already settled in Master
  Sheet 7: Summary            — financial and processing statistics

Column formatting matches Toll master data.xlsx exactly:
  - Unique Key:  formula =TRIM(D{r})&"|"&...
  - Dup Count:   formula =COUNTIF($A:$A,A{r})
  - TOTAL AMOUNT: formula =L{r}+M{r}
  - Toll Amount / Admin Fee: accounting currency format
  - TOTAL AMOUNT: red-negative currency format
  - Date (col D): long date format
  - Date Received / Hire Date / Return Date: string as stored
  - invoice added/sent/Follow-up: short date format when date object
"""

import io
from typing import Dict, Any, List, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .schemas import (
    MASTER_COLUMNS,
    SHEET_MASTER_DATA,
    SHEET_NEW_TOLLS,
    SHEET_CUSTOMER_SUMMARY,
    SHEET_DUPLICATES,
    SHEET_UNMATCHED,
    SHEET_ALREADY_PAID,
    SHEET_SUMMARY,
)

# ── Number formats matching Toll master data.xlsx ───────────────────────────
FMT_CURRENCY   = '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@'
FMT_TOTAL      = '"$"#,##0.00_);[Red]\\("$"#,##0.00\\)'
FMT_DATE_LONG  = '[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy'   # col D
FMT_DATE_SHORT = 'd-mmm-yy'                                # invoice/follow-up cols
FMT_GENERAL    = 'General'

# Column indices (1-based) for formatting
_L = MASTER_COLUMNS.index("Toll Amount")   + 1   # 12
_M = MASTER_COLUMNS.index("Admin Fee")     + 1   # 13
_N = MASTER_COLUMNS.index("TOTAL AMOUNT ") + 1   # 14
_D = MASTER_COLUMNS.index("Date")          + 1   # 4
_R = MASTER_COLUMNS.index("invoice added to 365") + 1           # 18
_S = MASTER_COLUMNS.index("Invoice\nsent to customer") + 1      # 19
_V = MASTER_COLUMNS.index("Follow-up Date") + 1                 # 22

# Approximate column widths (points) matching template
_COL_WIDTHS = [
    113, 13, 17, 44, 18, 13, 16, 17, 17, 68,
    37,  15, 14, 14, 43, 29, 19, 20, 19, 19,
    18,  15, 12,
]


def _write_master_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    master_df: pd.DataFrame,
    header_fill_hex: str = "1F4E78",
) -> None:
    """
    Write the Paste_eToll_Data sheet with exact formulas and number formats.
    master_df must already contain exactly the MASTER_COLUMNS columns.
    """
    # ── Header row ──────────────────────────────────────────────────────────
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid"
    )

    ws.append(MASTER_COLUMNS)
    ws.row_dimensions[1].height = 28
    for col_idx in range(1, len(MASTER_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    # ── Data rows ────────────────────────────────────────────────────────────
    for row_idx, row_tuple in enumerate(master_df.itertuples(index=False), start=2):
        row_vals = list(row_tuple)

        # Inject exact formulas
        row_vals[0]  = (
            f'=TRIM(D{row_idx})&"|"&TRIM(E{row_idx})&"|"&TRIM(H{row_idx})&"|"'
            f'&TRIM(I{row_idx})&"|"&TRIM(J{row_idx})&"|"'
            f'&SUBSTITUTE(TRIM(K{row_idx}),"-","")&"|"&TRIM(L{row_idx})'
        )
        row_vals[1]  = f'=COUNTIF($A:$A,A{row_idx})'
        row_vals[13] = f'=L{row_idx}+M{row_idx}'

        ws.append(row_vals)

        # Apply number formats
        ws.cell(row=row_idx, column=_L).number_format = FMT_CURRENCY
        ws.cell(row=row_idx, column=_M).number_format = FMT_CURRENCY
        ws.cell(row=row_idx, column=_N).number_format = FMT_TOTAL
        # Date col D — only format if it is actually a datetime object
        d_cell = ws.cell(row=row_idx, column=_D)
        if hasattr(d_cell.value, "year"):
            d_cell.number_format = FMT_DATE_LONG
        # Short-date columns
        for short_col in (_R, _S, _V):
            c = ws.cell(row=row_idx, column=short_col)
            if hasattr(c.value, "year"):
                c.number_format = FMT_DATE_SHORT

        ws.row_dimensions[row_idx].height = 18

    # ── Column widths ────────────────────────────────────────────────────────
    for i, w in enumerate(_COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Freeze top row ────────────────────────────────────────────────────────
    ws.freeze_panes = "A2"


# ── Generic audit-sheet helper ────────────────────────────────────────────────

def _write_audit_sheet(
    ws: openpyxl.worksheet.worksheet.Worksheet,
    headers: List[str],
    rows: List[List[Any]],
    header_fill_hex: str,
    currency_col_indices: Optional[List[int]] = None,
) -> None:
    """Write a simple audit table with coloured header."""
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(
        start_color=header_fill_hex, end_color=header_fill_hex, fill_type="solid"
    )
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    reg_font = Font(name="Calibri", size=10)

    ws.append(headers)
    ws.row_dimensions[1].height = 22
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    currency_set = set(currency_col_indices or [])

    for r_idx, row_data in enumerate(rows, start=2):
        ws.append(row_data)
        ws.row_dimensions[r_idx].height = 18
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = reg_font
            cell.border = border
            if c_idx in currency_set:
                cell.number_format = FMT_CURRENCY
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto column widths
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max(max_len + 3, 12), 50
        )


# ── Main public function ───────────────────────────────────────────────────────

def generate_excel_workbook(
    master_df: pd.DataFrame,
    new_tolls_df: pd.DataFrame,
    duplicates_df: pd.DataFrame,
    unmatched_df: pd.DataFrame,
    already_paid_df: pd.DataFrame,
    customer_summary_df: pd.DataFrame,
    summary_data: Dict[str, Any],
) -> bytes:
    """
    Generate the full multi-sheet output workbook in memory.
    Returns raw bytes suitable for writing to disk or streaming via HTTP.
    """
    wb = openpyxl.Workbook()
    if wb.active:
        wb.remove(wb.active)

    # ── Sheet 1: Paste_eToll_Data ─────────────────────────────────────────────
    ws_master = wb.create_sheet(title=SHEET_MASTER_DATA)
    export_master = master_df.copy()
    for col in MASTER_COLUMNS:
        if col not in export_master.columns:
            export_master[col] = ""
    export_master = export_master[MASTER_COLUMNS]
    _write_master_sheet(ws_master, export_master, header_fill_hex="1F4E78")

    # ── Sheet 2: New Tolls ────────────────────────────────────────────────────
    ws_new = wb.create_sheet(title=SHEET_NEW_TOLLS)
    export_new = new_tolls_df.copy() if not new_tolls_df.empty else pd.DataFrame(columns=MASTER_COLUMNS)
    for col in MASTER_COLUMNS:
        if col not in export_new.columns:
            export_new[col] = ""
    export_new = export_new[MASTER_COLUMNS]
    _write_master_sheet(ws_new, export_new, header_fill_hex="2E75B6")

    # ── Sheet 3: Customer Summary ─────────────────────────────────────────────
    ws_cust = wb.create_sheet(title=SHEET_CUSTOMER_SUMMARY)
    cust_headers = [
        "Driver", "REGO", "Toll Count",
        "Total Toll Amount", "Total Admin Fee", "Total Amount Payable"
    ]
    cust_rows: List[List[Any]] = []
    if not customer_summary_df.empty:
        for _, r in customer_summary_df.iterrows():
            cust_rows.append([
                r.get("Driver", ""),
                r.get("REGO", ""),
                int(r.get("Toll Count", 0)),
                float(r.get("Total Toll Amount", 0.0)),
                float(r.get("Total Admin Fee", 0.0)),
                float(r.get("Total Amount Payable", 0.0)),
            ])
    _write_audit_sheet(ws_cust, cust_headers, cust_rows, "1F4E78", currency_col_indices=[4, 5, 6])

    # ── Sheet 4: Duplicates ───────────────────────────────────────────────────
    ws_dup = wb.create_sheet(title=SHEET_DUPLICATES)
    dup_headers = ["REGO", "Toll Date", "Toll Time", "StartDateTime", "Toll Amount",
                   "Type", "Details", "Duplicate Key", "Reason"]
    dup_rows: List[List[Any]] = []
    if not duplicates_df.empty:
        for _, r in duplicates_df.iterrows():
            dup_rows.append([
                r.get("rego", ""), r.get("toll_date", ""), r.get("toll_time", ""),
                r.get("start_datetime", ""), r.get("toll_amount", 0.0),
                r.get("type", ""), r.get("details", ""),
                r.get("dup_key", ""), r.get("reason", "Duplicate"),
            ])
    _write_audit_sheet(ws_dup, dup_headers, dup_rows, "C00000", currency_col_indices=[5])

    # ── Sheet 5: Unmatched ────────────────────────────────────────────────────
    ws_unm = wb.create_sheet(title=SHEET_UNMATCHED)
    unm_headers = ["REGO", "Toll Date", "Toll Time", "StartDateTime", "Toll Amount",
                   "Type", "Details", "Status", "Reason"]
    unm_rows: List[List[Any]] = []
    if not unmatched_df.empty:
        for _, r in unmatched_df.iterrows():
            unm_rows.append([
                r.get("rego", ""), r.get("toll_date", ""), r.get("toll_time", ""),
                r.get("start_datetime", ""), r.get("toll_amount", 0.0),
                r.get("type", ""), r.get("details", ""),
                r.get("status", "UNMATCHED"), r.get("reason", ""),
            ])
    _write_audit_sheet(ws_unm, unm_headers, unm_rows, "ED7D31", currency_col_indices=[5])

    # ── Sheet 6: Already Paid ─────────────────────────────────────────────────
    ws_paid = wb.create_sheet(title=SHEET_ALREADY_PAID)
    paid_headers = ["REGO", "Toll Date", "Toll Time", "StartDateTime", "Toll Amount",
                    "Duplicate Key", "Reason"]
    paid_rows: List[List[Any]] = []
    if not already_paid_df.empty:
        for _, r in already_paid_df.iterrows():
            paid_rows.append([
                r.get("rego", ""), r.get("toll_date", ""), r.get("toll_time", ""),
                r.get("start_datetime", ""), r.get("toll_amount", 0.0),
                r.get("dup_key", ""), r.get("reason", "Already Paid"),
            ])
    _write_audit_sheet(ws_paid, paid_headers, paid_rows, "70AD47", currency_col_indices=[5])

    # ── Sheet 7: Summary ──────────────────────────────────────────────────────
    ws_sum = wb.create_sheet(title=SHEET_SUMMARY)

    title_cell = ws_sum.cell(row=2, column=2, value="TOLL PROCESSING AUTOMATION — SUMMARY")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="1F4E78")

    label_font = Font(name="Calibri", size=10, bold=True)
    value_font = Font(name="Calibri", size=10)

    meta = [
        ("Date Received",          summary_data.get("date_received", "")),
        ("Admin Fee per Toll",      f"${summary_data.get('admin_fee', 0):.2f}"),
        ("Report Generated",        pd.Timestamp.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (k, v) in enumerate(meta, start=4):
        ws_sum.cell(row=i, column=2, value=k).font = label_font
        ws_sum.cell(row=i, column=3, value=v).font = value_font

    processing = [
        ("", ""),
        ("=== EXISTING MASTER ===", ""),
        ("Rows before clean",       summary_data.get("master_rows_before_clean", 0)),
        ("Duplicate rows removed",  summary_data.get("master_dupes_removed", 0)),
        ("Rows after clean",        summary_data.get("master_rows_after_clean", 0)),
        ("", ""),
        ("=== LINKT PROCESSING ===", ""),
        ("Total Genuine LINKT Trips",    summary_data.get("total_linkt_records", 0)),
        ("  ├── Duplicates in Master (Total)", summary_data.get("total_duplicates_in_master", 0)),
        ("  │     ├── Unpaid Duplicates",      summary_data.get("duplicates_count", 0)),
        ("  │     └── Already Paid / Settled", summary_data.get("already_paid_count", 0)),
        ("  └── Genuinely New LINKT Records",  summary_data.get("new_tolls_count", 0) + summary_data.get("unmatched_count", 0)),
        ("        ├── Matched & Chargeable",   summary_data.get("new_tolls_count", 0)),
        ("        └── Unmatched (Manual Review)", summary_data.get("unmatched_count", 0)),
        ("", ""),
        ("=== FINANCIALS (NEW MATCHED TOLLS ONLY) ===", ""),
        ("Total Toll Amount",            f"${summary_data.get('total_toll_amount', 0):.2f}"),
        ("Total Admin Fees",             f"${summary_data.get('total_admin_fee', 0):.2f}"),
        ("Total Customer Payable",       f"${summary_data.get('total_customer_amount', 0):.2f}"),
        ("", ""),
        ("=== FINAL MASTER ===", ""),
        ("Final Master row count",  summary_data.get("final_master_row_count", 0)),
    ]
    start_row = 8
    for i, (k, v) in enumerate(processing, start=start_row):
        ws_sum.cell(row=i, column=2, value=k).font = label_font if k.startswith("=") else label_font
        ws_sum.cell(row=i, column=3, value=v).font = value_font
    ws_sum.column_dimensions["B"].width = 35
    ws_sum.column_dimensions["C"].width = 25

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
