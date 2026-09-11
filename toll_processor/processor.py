"""Core pipeline coordinator for Toll Processing Automation.

Workflow:
  1. Load existing Master (from template or stored current_master.xlsx).
  2. Clean existing Master: remove duplicate rows using completeness-based
     selection; merge useful metadata; preserve payment/invoice info.
  3. Read LINKT file; filter genuine vehicle Trip records only.
  4. For each genuine LINKT toll:
       a. Build dup_key: Normalized REGO + StartDateTime + Positive Toll Amount.
       b. Check against cleaned Master index -> DUPLICATE / ALREADY PAID.
       c. Check against current batch -> DUPLICATE.
       d. Match driver from 365 Booking Export.
       e. If matched: calculate Toll Amount, Admin Fee $5.55, TOTAL AMOUNT.
       f. If unmatched/multiple-match: report only, do NOT add to Master.
  5. Build final Master = cleaned historical rows + new matched tolls only.
  6. Final duplicate audit on entire Master to catch any edge cases.
  7. Generate multi-sheet Excel workbook.
"""

from dataclasses import dataclass, field
from typing import Dict, Any, List, Optional, Tuple, Union, BinaryIO
import pandas as pd

from .schemas import (
    MASTER_COLUMNS,
    STATUS_NEW,
    STATUS_DUPLICATE,
    STATUS_ALREADY_PAID,
    STATUS_UNMATCHED,
    STATUS_MULTIPLE_MATCH,
    DEFAULT_ADMIN_FEE,
    DEFAULT_DATE_RECEIVED,
)
from .normalizer import parse_flexible_datetime, normalize_rego
from .file_readers import read_linkt_file, read_365_booking_file, read_master_file
from .matcher import match_driver_for_toll
from .duplicate_detector import DuplicateDetector, clean_master_dataframe
from .fee_calculator import calculate_toll_totals
from .excel_generator import generate_excel_workbook
from . import master_store


@dataclass
class ProcessingResult:
    """Encapsulates the full output of a toll processing batch."""
    master_df: pd.DataFrame
    new_tolls_df: pd.DataFrame
    duplicates_df: pd.DataFrame
    unmatched_df: pd.DataFrame
    already_paid_df: pd.DataFrame
    customer_summary_df: pd.DataFrame
    master_cleaned_dupes: List[Dict[str, Any]]  # rows purged from existing Master
    summary_stats: Dict[str, Any]
    excel_bytes: bytes
    detected_linkt_cols: Dict[str, str] = field(default_factory=dict)
    detected_365_cols: Dict[str, str] = field(default_factory=dict)


def process_toll_files(
    linkt_file: Union[str, bytes, BinaryIO],
    bookings_365_file: Union[str, bytes, BinaryIO],
    master_file: Optional[Union[str, bytes, BinaryIO]] = None,
    date_received: Union[str, pd.Timestamp] = DEFAULT_DATE_RECEIVED,
    admin_fee: float = DEFAULT_ADMIN_FEE,
) -> ProcessingResult:
    """
    Main processing function.

    master_file: path/bytes/stream to the Master Excel (Toll master data.xlsx or
                 current_master.xlsx). If None, starts with an empty Master.
    """
    # --- 1. Parse date_received ---
    parsed_date_rec = parse_flexible_datetime(date_received)
    if parsed_date_rec is None:
        parsed_date_rec = pd.Timestamp.now()
    formatted_date_received = parsed_date_rec.strftime("%d-%b-%Y")   # e.g. 08-Sep-2026
    file_date_suffix = parsed_date_rec.strftime("%Y-%m-%d")

    # --- 2. Read input files ---
    linkt_df, linkt_mapping = read_linkt_file(linkt_file)
    bookings_df, bookings_mapping = read_365_booking_file(bookings_365_file)

    if master_file is not None:
        master_raw_df, _ = read_master_file(master_file)
    else:
        master_raw_df = pd.DataFrame(columns=MASTER_COLUMNS)

    # --- 3. Clean existing Master (remove duplicates, keep best rows) ---
    master_cleaned_df, master_purged = clean_master_dataframe(master_raw_df)

    rows_before_clean = len(master_raw_df)
    rows_after_clean  = len(master_cleaned_df)
    master_dupes_removed = rows_before_clean - rows_after_clean

    # --- 4. Index cleaned Master for LINKT duplicate detection ---
    dup_detector = DuplicateDetector(master_df=master_cleaned_df)

    # --- 5. Process each genuine LINKT toll ---
    new_tolls_rows:    List[Dict[str, Any]] = []
    duplicates_rows:   List[Dict[str, Any]] = []
    unmatched_rows:    List[Dict[str, Any]] = []
    already_paid_rows: List[Dict[str, Any]] = []

    for idx, row in linkt_df.iterrows():
        raw_rego   = row.get("raw_rego", "")
        start_dt   = row.get("parsed_start_dt")
        end_dt     = row.get("parsed_end_dt")
        toll_amt   = row.get("toll_amount", 0.0)
        dup_key    = row.get("dup_key", "")
        toll_type  = row.get("type", "Trips")
        details    = row.get("details", "")
        concession = row.get("concession", "")

        # Formatted strings
        date_str     = start_dt.strftime("%d/%m/%Y") if start_dt is not None else ""
        time_str     = start_dt.strftime("%H:%M") if start_dt is not None else ""
        start_dt_str = start_dt.strftime("%d/%m/%Y %H:%M") if start_dt is not None else ""
        end_dt_str   = end_dt.strftime("%d/%m/%Y %H:%M") if end_dt is not None else start_dt_str

        # Step A: Duplicate / Paid check
        is_dup, dup_status, dup_reason = dup_detector.check_duplicate(dup_key)
        if is_dup:
            record = {
                "rego": raw_rego,
                "toll_date": date_str,
                "toll_time": time_str,
                "start_datetime": start_dt_str,
                "toll_amount": toll_amt,
                "type": toll_type,
                "details": details,
                "dup_key": dup_key,
                "reason": dup_reason,
            }
            if dup_status == STATUS_ALREADY_PAID:
                already_paid_rows.append(record)
            else:
                duplicates_rows.append(record)
            continue

        # Step B: Driver matching from 365
        match_result = match_driver_for_toll(row, bookings_df)

        if not match_result["matched"]:
            unmatched_rows.append({
                "rego": raw_rego,
                "toll_date": date_str,
                "toll_time": time_str,
                "start_datetime": start_dt_str,
                "toll_amount": toll_amt,
                "type": toll_type,
                "details": details,
                "status": match_result.get("status", STATUS_UNMATCHED),
                "reason": match_result.get("reason", "No match found"),
            })
            continue

        # Step C: Matched — calculate amounts
        pos_toll, calc_fee, total_charge = calculate_toll_totals(toll_amt, admin_fee)

        hire_dt   = match_result.get("hire_date")
        return_dt = match_result.get("return_date")

        # Format Hire Date  (preserve 365 text format when possible)
        if isinstance(hire_dt, pd.Timestamp):
            hire_dt_str = hire_dt.strftime("%d %b %Y %H:%M")
        else:
            hire_dt_str = str(hire_dt or "")

        # Format Return Date
        if isinstance(return_dt, pd.Timestamp):
            return_dt_str = return_dt.strftime("%d %b %Y %H:%M")
        elif return_dt == "RETURN_NOT_SET":
            return_dt_str = "Return not set"
        else:
            return_dt_str = str(return_dt or "")

        master_row: Dict[str, Any] = {
            # Cols 1-2 are formula-driven in Excel; set placeholders
            "Unique Key":                  "",
            "Duplicate Count":             "",
            "Date Received":               formatted_date_received,
            "Date":                        start_dt if start_dt is not None else "",
            "Time":                        time_str,
            "StartDateTime":               start_dt_str,
            "EndDateTime":                 end_dt_str,
            "Type of activity":            toll_type,
            "REGO":                        raw_rego,
            "Details":                     details,
            "Concession":                  concession,
            "Toll Amount":                 pos_toll,
            "Admin Fee":                   calc_fee,
            "TOTAL AMOUNT ":              total_charge,
            "Driver":                      match_result.get("driver", ""),
            "Hire Date":                   hire_dt_str,
            "Return Date":                 return_dt_str,
            "invoice added to 365":        "No",
            "Invoice\nsent to customer":   "No",
            "Payment Method":              "",
            "Payment Status":              "Pending",
            "Follow-up Date":              "",
            "Source":                      "Linkt",
        }
        new_tolls_rows.append(master_row)

    # --- 6. Build DataFrames ---
    new_tolls_df    = pd.DataFrame(new_tolls_rows, columns=MASTER_COLUMNS) if new_tolls_rows else pd.DataFrame(columns=MASTER_COLUMNS)
    duplicates_df   = pd.DataFrame(duplicates_rows)   if duplicates_rows   else pd.DataFrame()
    unmatched_df    = pd.DataFrame(unmatched_rows)    if unmatched_rows    else pd.DataFrame()
    already_paid_df = pd.DataFrame(already_paid_rows) if already_paid_rows else pd.DataFrame()

    # --- 7. Build updated Master (clean historical + new matched tolls) ---
    # Ensure master_cleaned_df has the same MASTER_COLUMNS structure
    for col in MASTER_COLUMNS:
        if col not in master_cleaned_df.columns:
            master_cleaned_df[col] = ""
    master_cleaned_df = master_cleaned_df[MASTER_COLUMNS]

    updated_master_df = pd.concat(
        [master_cleaned_df, new_tolls_df], ignore_index=True
    )

    # --- 8. Final duplicate audit (catch any edge cases) ---
    updated_master_df, final_purged = clean_master_dataframe(updated_master_df)
    if final_purged:
        # These shouldn't happen in normal operation — log them as duplicates
        extra_dup_rows = [
            {
                "rego": r.get("REGO", ""),
                "toll_date": str(r.get("StartDateTime", ""))[:10],
                "toll_time": str(r.get("Time", "")),
                "start_datetime": str(r.get("StartDateTime", "")),
                "toll_amount": r.get("Toll Amount", 0.0),
                "dup_key": "",
                "reason": "Caught by final duplicate audit",
            }
            for r in final_purged
        ]
        duplicates_df = pd.concat(
            [duplicates_df, pd.DataFrame(extra_dup_rows)], ignore_index=True
        ) if not duplicates_df.empty else pd.DataFrame(extra_dup_rows)

    # --- 9. Customer Summary ---
    if not new_tolls_df.empty:
        grp = (
            new_tolls_df
            .groupby(["Driver", "REGO"])
            .agg(
                Toll_Count        =("Toll Amount",   "count"),
                Total_Toll_Amount =("Toll Amount",   "sum"),
                Total_Admin_Fee   =("Admin Fee",     "sum"),
                Total_Amt_Payable =("TOTAL AMOUNT ", "sum"),
            )
            .reset_index()
        )
        grp.columns = ["Driver", "REGO", "Toll Count",
                       "Total Toll Amount", "Total Admin Fee", "Total Amount Payable"]
        grp["Total Toll Amount"]   = grp["Total Toll Amount"].round(2)
        grp["Total Admin Fee"]     = grp["Total Admin Fee"].round(2)
        grp["Total Amount Payable"] = grp["Total Amount Payable"].round(2)
        customer_summary_df = grp
    else:
        customer_summary_df = pd.DataFrame(
            columns=["Driver", "REGO", "Toll Count",
                     "Total Toll Amount", "Total Admin Fee", "Total Amount Payable"]
        )

    # --- 10. Summary statistics ---
    total_toll_amt  = float(new_tolls_df["Toll Amount"].sum())   if not new_tolls_df.empty else 0.0
    total_fee_amt   = float(new_tolls_df["Admin Fee"].sum())     if not new_tolls_df.empty else 0.0
    total_cust_amt  = float(new_tolls_df["TOTAL AMOUNT "].sum()) if not new_tolls_df.empty else 0.0

    unmatched_count      = len(unmatched_df[unmatched_df["status"] == STATUS_UNMATCHED])      if not unmatched_df.empty else 0
    multi_match_count    = len(unmatched_df[unmatched_df["status"] == STATUS_MULTIPLE_MATCH]) if not unmatched_df.empty else 0

    duplicates_count_val = len(duplicates_df) if not duplicates_df.empty else 0
    already_paid_count_val = len(already_paid_df) if not already_paid_df.empty else 0
    total_dupes_in_master_val = duplicates_count_val + already_paid_count_val

    summary_stats: Dict[str, Any] = {
        "date_received":               formatted_date_received,
        "file_date_suffix":            file_date_suffix,
        "admin_fee":                   admin_fee,
        "master_rows_before_clean":    rows_before_clean,
        "master_dupes_removed":        master_dupes_removed,
        "master_rows_after_clean":     rows_after_clean,
        "total_records":               len(linkt_df),
        "total_linkt_records":         len(linkt_df),
        "new_tolls_count":             len(new_tolls_rows),
        "duplicates_count":            duplicates_count_val,
        "already_paid_count":          already_paid_count_val,
        "total_duplicates_in_master":  total_dupes_in_master_val,
        "unmatched_count":             unmatched_count,
        "multi_match_count":           multi_match_count,
        "total_unmatched_all":         len(unmatched_df) if not unmatched_df.empty else 0,
        "total_toll_amount":           round(total_toll_amt, 2),
        "total_admin_fee":             round(total_fee_amt, 2),
        "total_customer_amount":       round(total_cust_amt, 2),
        "final_master_row_count":      len(updated_master_df),
    }

    # --- 11. Generate full Excel workbook ---
    excel_bytes = generate_excel_workbook(
        master_df           = updated_master_df,
        new_tolls_df        = new_tolls_df,
        duplicates_df       = duplicates_df,
        unmatched_df        = unmatched_df,
        already_paid_df     = already_paid_df,
        customer_summary_df = customer_summary_df,
        summary_data        = summary_stats,
    )

    return ProcessingResult(
        master_df           = updated_master_df,
        new_tolls_df        = new_tolls_df,
        duplicates_df       = duplicates_df,
        unmatched_df        = unmatched_df,
        already_paid_df     = already_paid_df,
        customer_summary_df = customer_summary_df,
        master_cleaned_dupes= master_purged,
        summary_stats       = summary_stats,
        excel_bytes         = excel_bytes,
        detected_linkt_cols = linkt_mapping,
        detected_365_cols   = bookings_mapping,
    )


def process_daily_run(
    linkt_file: Union[str, bytes, BinaryIO],
    bookings_365_file: Union[str, bytes, BinaryIO],
    master_file: Optional[Union[str, bytes, BinaryIO]] = None,
    date_received: Union[str, pd.Timestamp] = DEFAULT_DATE_RECEIVED,
    admin_fee: float = DEFAULT_ADMIN_FEE,
    is_initial_setup: bool = False,
) -> Tuple["ProcessingResult", Dict[str, Any]]:
    """
    Daily processing wrapper with automatic Master carry-forward and archiving.

    Initial Setup (is_initial_setup=True):
      - master_file must be provided (Toll master data.xlsx).
      - Saves result as data/current_master.xlsx.

    Daily Run (is_initial_setup=False):
      - If master_file is None, auto-loads data/current_master.xlsx.
      - Archives old master, saves new one.
    """
    if not is_initial_setup and master_file is None:
        _, stored_path = master_store.load_stored_master()
        master_file = stored_path
    elif master_file is not None and hasattr(master_file, "seek"):
        pass  # stream already positioned

    result = process_toll_files(
        linkt_file          = linkt_file,
        bookings_365_file   = bookings_365_file,
        master_file         = master_file,
        date_received       = date_received,
        admin_fee           = admin_fee,
    )

    # Generate master-only bytes for persistence (clean workbook, Paste_eToll_Data only)
    master_only_bytes = _generate_master_only_excel(result.master_df)

    save_info = master_store.save_new_master(
        excel_bytes = master_only_bytes,
        date_suffix = result.summary_stats.get("file_date_suffix", ""),
    )

    return result, save_info


def _generate_master_only_excel(master_df: pd.DataFrame) -> bytes:
    """Generate a clean Excel file containing only the Paste_eToll_Data sheet."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Paste_eToll_Data"

    export_df = master_df.copy()
    for col in MASTER_COLUMNS:
        if col not in export_df.columns:
            export_df[col] = ""
    export_df = export_df[MASTER_COLUMNS]

    # Header row
    ws.append(MASTER_COLUMNS)
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col_idx in range(1, len(MASTER_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rows with exact formulas
    for row_idx, row_tuple in enumerate(export_df.itertuples(index=False), start=2):
        row_vals = list(row_tuple)
        # Col A: Unique Key formula
        row_vals[0] = (
            f'=TRIM(D{row_idx})&"|"&TRIM(E{row_idx})&"|"&TRIM(H{row_idx})&"|"'
            f'&TRIM(I{row_idx})&"|"&TRIM(J{row_idx})&"|"'
            f'&SUBSTITUTE(TRIM(K{row_idx}),"-","")&"|"&TRIM(L{row_idx})'
        )
        # Col B: Duplicate Count formula
        row_vals[1] = f'=COUNTIF($A:$A,A{row_idx})'
        # Col N: TOTAL AMOUNT formula
        row_vals[13] = f'=L{row_idx}+M{row_idx}'
        ws.append(row_vals)

    # Apply number formats
    from openpyxl.utils import get_column_letter
    currency_fmt = '_-"$"* #,##0.00_-;\\-"$"* #,##0.00_-;_-"$"* "-"??_-;_-@'
    total_fmt    = '"$"#,##0.00_);[Red]\\("$"#,##0.00\\)'
    date_full_fmt = '[$-F800]dddd\\,\\ mmmm\\ dd\\,\\ yyyy'
    date_short_fmt = 'd-mmm-yy'

    currency_cols = [12, 13]  # L, M (1-based)
    total_col     = 14         # N
    date_full_col = 4          # D
    date_short_cols = [18, 19, 22]  # R, S, V

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in currency_cols:
            ws.cell(row=row_idx, column=col_idx).number_format = currency_fmt
        ws.cell(row=row_idx, column=total_col).number_format = total_fmt
        ws.cell(row=row_idx, column=date_full_col).number_format = date_full_fmt
        for col_idx in date_short_cols:
            ws.cell(row=row_idx, column=col_idx).number_format = date_short_fmt

    # Freeze top row
    ws.freeze_panes = "A2"

    # Column widths matching Toll master data.xlsx approximately
    col_widths = [
        113, 13, 17, 44, 18, 13, 16, 17, 17, 68,
        37, 15, 14, 14, 43, 29, 19, 20, 19, 19,
        18, 15, 12
    ]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(w, 113)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
